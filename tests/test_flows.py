"""شبیه‌سازی end-to-end جریان‌های بات با mock کردن Message/CallbackQuery/FSMContext.

هیچ ارتباطی با تلگرام برقرار نمی‌شود و توکن واقعی لازم نیست.
"""
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import jdatetime
import pytest
from openpyxl import load_workbook

from app import bot as B
from app import calendar_utils as cal
from app import reminders as rem
from app.config import settings, BOT_NAME
from app.db import connect, get_user, memberships_of, active_membership, calendar_rules

# asyncio_mode=auto در pytest.ini؛ تست‌های async بدون mark اجرا می‌شوند.


# ---------- تدارکات mock ----------

class FakeState:
    def __init__(self):
        self.state = None
        self.data = {}

    async def set_state(self, s):
        self.state = s

    async def update_data(self, **kw):
        self.data.update(kw)
        return dict(self.data)

    async def get_data(self):
        return dict(self.data)

    async def clear(self):
        self.state = None
        self.data = {}


def msg(text, uid=5001, name="کاربر تست", username="tester", contact=None):
    m = MagicMock()
    m.text = text
    m.contact = contact
    m.from_user = MagicMock(id=uid, full_name=name, username=username)
    m.answer = AsyncMock()
    m.answer_document = AsyncMock()
    return m


def cbq(data, uid=5001):
    c = MagicMock()
    c.data = data
    c.from_user = MagicMock(id=uid, full_name="کاربر تست", username="tester")
    c.message = msg(None, uid)
    c.answer = AsyncMock()
    return c


def said(mock_msg):
    return "\n".join(str(c.args[0]) for c in mock_msg.answer.call_args_list if c.args)


@pytest.fixture(autouse=True)
def capture_admin_notifications(monkeypatch):
    sent = []
    monkeypatch.setattr(B.bot, "send_message", AsyncMock(side_effect=lambda cid, t, **k: sent.append((cid, t))))
    return sent


ROLE_BTN = {"generator": "⚡️ نیروگاه", "buyer": "🏭 خریدار نهایی / صنعت",
            "retailer": "🏬 خرده‌فروش", "broker": "🤝 کارگزاری"}


async def register(uid, role, org_name, capacity="5000", phone="09121234567", st=None):
    st = st or FakeState()
    transcript, last = [], None
    steps = [(ROLE_BTN[role], B.reg_role), (org_name, B.reg_org_name), ("علی رضایی", B.reg_person),
             ("مدیرعامل", B.reg_position), (phone, B.reg_phone)]
    if role in B.CAPACITY_ROLES:
        steps.append((capacity, B.reg_capacity))
    for text, handler in steps:
        last = msg(text, uid)
        await handler(last, st)
        transcript.append(said(last))
    return last, "\n".join(transcript)


async def do_trade(flow, uid, ptype="weekly", volume=5000, price=80000, period_index=0):
    st = FakeState()
    m = msg(B.BTN_OFFER if flow == "offer" else B.BTN_REQUEST, uid)
    await B.start_trade(flow, m, st)
    c = cbq(f"pt:{flow}:{ptype}", uid)
    await B.choose_period_type(c, st)
    kb = c.message.answer.call_args_list[-1].kwargs["reply_markup"]
    pick = cbq(f"pd:{flow}:{period_index}", uid)
    await B.choose_period(pick, st)
    vol = msg(str(volume), uid)
    await B.trade_volume(vol, st)
    final = msg(str(price), uid)
    await B.trade_price(final, st)
    return final, kb, said(pick.message) + said(vol)


# ---------- برند و «درباره ما» ----------

async def test_about_uses_the_new_text_without_company_name(fresh_db):
    m = msg("/about")
    await B.about(m)
    out = said(m)
    assert BOT_NAME in out and "سامانه تطبیق و تحلیل بازار برق" in out
    assert "محرمانه" in out and "@Farbod_Papoli" in out and "داتیس آوان نیرو" in out
    assert not hasattr(settings, "company_name")


async def test_bot_name_is_used_in_panels(fresh_db):
    m = msg("/admin", settings.admin_chat_ids[0])
    await B.admin_panel(m)
    assert BOT_NAME in said(m)


async def test_invite_button_shares_bot_link(fresh_db, monkeypatch):
    B._link_cache.clear()
    monkeypatch.setattr(settings, "bot_username", "PowerMatchBot")
    m = msg(B.BTN_INVITE)
    await B.invite(m)
    assert "https://t.me/PowerMatchBot" in said(m)
    kb = m.answer.call_args_list[-1].kwargs["reply_markup"]
    assert kb.inline_keyboard[0][0].switch_inline_query is not None


# ---------- اعتبارسنجی شماره تلفن ----------

@pytest.mark.parametrize("raw,expected", [
    ("09121234567", "09121234567"), ("۰۹۱۲۱۲۳۴۵۶۷", "09121234567"),
    ("+989121234567", "09121234567"), ("00989121234567", "09121234567"),
    ("9121234567", "09121234567"), ("0912 123 4567", "09121234567"),
    ("0912-123-4567", "09121234567"),
])
def test_valid_iranian_mobiles(raw, expected):
    assert B.normalize_phone(raw) == expected


@pytest.mark.parametrize("raw", ["0912123456", "091212345678", "08121234567",
                                 "02112345678", "abcd", "", "+15551234567"])
def test_invalid_mobiles_are_rejected(raw):
    assert B.normalize_phone(raw) is None


async def test_registration_rejects_bad_phone_and_asks_again(fresh_db):
    st = FakeState()
    await B.reg_role(msg(ROLE_BTN["generator"]), st)
    await B.reg_org_name(msg("آرمانی انرژی آسیا"), st)
    await B.reg_person(msg("علی رضایی"), st)
    await B.reg_position(msg("مدیرعامل"), st)
    bad = msg("0912123456")
    await B.reg_phone(bad, st)
    assert "نامعتبر" in said(bad) and "09121234567" in said(bad)
    assert st.state == B.Registration.phone            # به مرحله‌ی بعد نمی‌رود


async def test_phone_can_be_shared_via_contact_button(fresh_db):
    st = FakeState()
    await B.reg_role(msg(ROLE_BTN["retailer"]), st)
    await B.reg_org_name(msg("خرده‌فروش نمونه"), st)
    await B.reg_person(msg("علی رضایی"), st)
    prompt = msg("مدیرعامل")
    await B.reg_position(prompt, st)
    kb = prompt.answer.call_args_list[-1].kwargs["reply_markup"]
    assert kb.keyboard[0][0].request_contact is True   # گزینه‌ی اختیاری، نه سؤال جداگانه

    contact = MagicMock(phone_number="+989121234567", user_id=5001)
    await B.reg_phone_contact(msg(None, 5001, contact=contact), st)
    org = memberships_of(get_user(5001)["id"])[0]
    assert org["phone"] == "09121234567"


async def test_contact_of_another_account_is_rejected(fresh_db):
    st = FakeState()
    await B.reg_role(msg(ROLE_BTN["retailer"]), st)
    await B.reg_org_name(msg("خرده‌فروش نمونه"), st)
    await B.reg_person(msg("علی رضایی"), st)
    await B.reg_position(msg("مدیرعامل"), st)
    foreign = msg(None, 5001, contact=MagicMock(phone_number="+989121234567", user_id=9999))
    await B.reg_phone_contact(foreign, st)
    assert "خودتان" in said(foreign)


# ---------- قالب‌بندی ----------

async def test_user_confirmation_is_bold_and_line_separated(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    final, _, _ = await do_trade("offer", 5001)
    out = said(final)
    assert "<b>✅ عرضه شما ثبت شد</b>" in out
    assert "<b>حجم:</b>" in out and "<b>کف قیمت:</b>" in out
    assert out.count("\n") >= 5                        # هر مورد در یک خط جدا


async def test_admin_order_report_uses_the_requested_layout(fresh_db, capture_admin_notifications):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await do_trade("offer", 5001, volume=5000, price=80000)
    note = [t for _, t in capture_admin_notifications if "اطلاعات عرضه" in t][0]
    for label in ("نام شرکت", "نام شخص", "شماره تلفن", "تاریخ", "حجم", "کف قیمت"):
        assert f"<b>{label}:</b>" in note
    assert "سقف قیمت" not in note                      # فقط عبارت متناسب با نوع سفارش
    assert "+989121234567" in note and "علی رضایی" in note   # قابل کلیک برای تماس


async def test_request_report_shows_ceiling_price_wording(fresh_db, capture_admin_notifications):
    await register(5002, "buyer", "فولاد نمونه", "3000")
    await do_trade("request", 5002)
    note = [t for _, t in capture_admin_notifications if "اطلاعات تقاضا" in t][0]
    assert "<b>سقف قیمت:</b>" in note and "کف قیمت" not in note


async def test_profile_is_shown_one_item_per_line_with_edit_buttons(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    m = msg(B.BTN_PROFILE)
    await B.profile(m)
    out = said(m)
    for label in ("نقش", "نام شخص مسئول", "سمت", "شماره تلفن", "سقف مجاز"):
        assert f"<b>{label}:</b>" in out
    kb = m.answer.call_args_list[-1].kwargs["reply_markup"]
    datas = [x[0].callback_data for x in kb.inline_keyboard]
    assert any(d.endswith(":phone") for d in datas) and any(d.startswith("pdel:") for d in datas)


async def test_user_input_cannot_inject_html(fresh_db):
    await register(5001, "generator", "<b>هک</b><script>x</script>")
    m = msg(B.BTN_PROFILE)
    await B.profile(m)
    out = said(m)
    assert "<script>" not in out and "&lt;script&gt;" in out


# ---------- ویرایش پروفایل ----------

async def test_user_can_edit_own_profile_phone(fresh_db, capture_admin_notifications):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    mid = memberships_of(get_user(5001)["id"])[0]["membership_id"]
    st = FakeState()
    await B.profile_edit_start(cbq(f"pedit:{mid}:phone"), st)
    assert st.state == B.EditProfile.value

    bad = msg("123")
    await B.profile_edit_save(bad, st)
    assert "نامعتبر" in said(bad)

    ok = msg("09351112233")
    await B.profile_edit_save(ok, st)
    assert memberships_of(get_user(5001)["id"])[0]["phone"] == "09351112233"
    assert "به‌روزرسانی شد" in said(ok)
    assert any("ویرایش پروفایل" in t for _, t in capture_admin_notifications)

    conn = connect()
    a = conn.execute("SELECT * FROM audit_log WHERE action='update'").fetchone()
    conn.close()
    assert a["old_value"] == "09121234567" and a["new_value"] == "09351112233"
    assert a["created_at_jalali"]


async def test_user_cannot_edit_another_users_profile(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await register(5002, "buyer", "فولاد نمونه", "3000")
    victim = memberships_of(get_user(5002)["id"])[0]["membership_id"]
    c = cbq(f"pedit:{victim}:phone", uid=5001)
    st = FakeState()
    await B.profile_edit_start(c, st)
    c.answer.assert_awaited_with("پروفایل یافت نشد.", show_alert=True)


# ---------- حذف پروفایل ----------

async def test_user_can_delete_a_whole_profile_but_history_survives(fresh_db):
    await register(5001, "generator", "نیروگاه اول")
    await do_trade("offer", 5001)
    await register(5001, "buyer", "فولاد دوم", "3000")
    u = get_user(5001)
    gen = next(r for r in memberships_of(u["id"]) if r["role"] == "generator")

    await B.profile_delete(cbq(f"pdelok:{gen['membership_id']}"))
    left = memberships_of(u["id"])
    assert [r["name"] for r in left] == ["فولاد دوم"]
    assert active_membership(u["id"])["name"] == "فولاد دوم"

    conn = connect()
    assert conn.execute("SELECT COUNT(*) c FROM organizations").fetchone()["c"] == 2   # حذف نرم
    o = conn.execute("SELECT * FROM offers").fetchone()
    assert o["deleted_at"] and o["status"] == "deleted"
    assert conn.execute("SELECT COUNT(*) c FROM audit_log WHERE action='delete'").fetchone()["c"] >= 2
    conn.close()

    m = msg(B.BTN_ORDERS, 5001)
    await B.my_orders(m)
    assert "هنوز سفارشی ثبت نکرده‌اید" in said(m)


# ---------- ویرایش و حذف سفارش ----------

async def test_user_can_edit_order_volume_and_price(fresh_db, capture_admin_notifications):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await do_trade("offer", 5001, volume=5000, price=80000)
    st = FakeState()
    await B.order_edit_start(cbq("oedit:o:1:volume"), st)
    await B.order_edit_save(msg("7500"), st)
    await B.order_edit_start(cbq("oedit:o:1:price"), st)
    ok = msg("91000")
    await B.order_edit_save(ok, st)

    conn = connect()
    row = conn.execute("SELECT * FROM offers WHERE id=1").fetchone()
    trail = conn.execute("SELECT * FROM audit_log WHERE entity='offer' AND action='update'").fetchall()
    conn.close()
    assert row["volume_kw"] == 7500 and row["min_price"] == 91000 and row["updated_at"]
    assert {t["field"] for t in trail} == {"حجم", "کف قیمت"}
    assert "به‌روزرسانی شد" in said(ok)
    assert any("ویرایش عرضه" in t and "+989121234567" in t for _, t in capture_admin_notifications)


async def test_user_can_change_order_period(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await do_trade("offer", 5001)
    options = cal.available_periods("weekly", calendar_rules())
    st = FakeState()
    await B.order_edit_start(cbq("oedit:o:1:period"), st)
    await B.order_edit_period_type(cbq("opt:o:1:weekly"), st)
    await B.order_edit_period_save(cbq("opd:o:1:2"), st)
    conn = connect()
    row = conn.execute("SELECT * FROM offers WHERE id=1").fetchone()
    conn.close()
    assert row["delivery_period"] == options[2][0]


async def test_user_can_delete_own_order_only(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await do_trade("offer", 5001)
    await register(5002, "generator", "نیروگاه دیگر")
    await do_trade("offer", 5002)

    intruder = cbq("odelok:o:2", uid=5001)             # سفارش کاربر دیگر
    await B.order_delete(intruder)
    intruder.answer.assert_awaited_with("سفارش یافت نشد.", show_alert=True)

    await B.order_delete(cbq("odelok:o:1", uid=5001))
    conn = connect()
    rows = {r["id"]: r["deleted_at"] for r in conn.execute("SELECT id,deleted_at FROM offers")}
    conn.close()
    assert rows[1] is not None and rows[2] is None


# ---------- یادآوری‌ها ----------

async def test_reminder_targets_only_users_missing_the_nearest_periods(fresh_db):
    await register(5001, "generator", "بدون سفارش")
    await register(5002, "generator", "با سفارش کامل", st=FakeState())
    await do_trade("offer", 5002, "weekly")
    await do_trade("offer", 5002, "monthly")

    pending = {name for _, name, _, _ in rem.pending_reminders()}
    assert "بدون سفارش" in pending and "با سفارش کامل" not in pending


async def test_reminder_lists_only_the_missing_period_type(fresh_db):
    await register(5001, "generator", "نیمه‌کاره")
    await do_trade("offer", 5001, "weekly")
    kinds = [k for _, name, _, miss in rem.pending_reminders() if name == "نیمه‌کاره" for k, _ in miss]
    assert kinds == ["ماهانه"]


async def test_reminder_audience_can_be_limited_by_role(fresh_db):
    await register(5001, "generator", "نیروگاه یک")
    await register(5002, "broker", "کارگزاری یک")
    rem.set_roles(["broker"])
    names = {name for _, name, _, _ in rem.pending_reminders()}
    assert names == {"کارگزاری یک"}
    rem.set_roles(list(rem.ALL_ROLES))
    assert len(rem.pending_reminders()) == 2


async def test_reminders_are_sent_only_at_configured_times_and_once(fresh_db, capture_admin_notifications):
    from datetime import datetime
    await register(5001, "generator", "بدون سفارش")
    rem.set_times("09:30,11:00")
    day = datetime(2026, 8, 14, 9, 30)
    assert await rem.send_due(B.bot, now=day.replace(hour=9, minute=0)) == 0     # ساعت غیرمجاز
    assert await rem.send_due(B.bot, now=day) == 1
    assert await rem.send_due(B.bot, now=day) == 0                               # تکرار نمی‌شود
    assert await rem.send_due(B.bot, now=day.replace(hour=11, minute=0)) == 1    # نوبت دوم
    assert "یادآوری ثبت سفارش" in capture_admin_notifications[-1][1]
    assert "بدون سفارش" in capture_admin_notifications[-1][1]


async def test_disabled_reminders_send_nothing(fresh_db):
    from app.db import set_setting
    from datetime import datetime
    await register(5001, "generator", "بدون سفارش")
    set_setting(rem.K_ENABLED, "0")
    assert await rem.send_due(B.bot, now=datetime(2026, 8, 14, 9, 30)) == 0


@pytest.mark.parametrize("raw,expected", [
    ("09:30,11:00", ["09:30", "11:00"]), ("9:5", ["09:05"]),
    ("11:00, 09:30", ["09:30", "11:00"]), ("۰۹:۳۰".translate(B.DIGITS), ["09:30"]),
])
def test_time_parsing(raw, expected):
    assert rem.parse_times(raw) == expected


@pytest.mark.parametrize("raw", ["25:00", "09:70", "abc", "", "09"])
def test_invalid_times_rejected(raw):
    with pytest.raises(ValueError):
        rem.parse_times(raw)


async def test_admin_can_configure_reminders_from_panel(fresh_db):
    admin = settings.admin_chat_ids[0]
    st = FakeState()
    c = cbq("ar:times", admin)
    await B.admin_reminders(c, st)
    assert st.state == B.AdminForm.rem_times

    bad = msg("99:99", admin)
    await B.save_rem_times(bad, st)
    assert "نامعتبر" in said(bad)
    await B.save_rem_times(msg("08:00,14:30", admin), st)
    assert rem.get_config()["times"] == ["08:00", "14:30"]

    await B.admin_reminders(cbq("ar:toggle", admin), st)
    assert rem.get_config()["enabled"] is False

    await B.admin_reminders(cbq("ar:role:all", admin), st)
    assert rem.get_config()["all_roles"] is True
    await B.admin_reminders(cbq("ar:role:generator", admin), st)
    assert set(rem.get_config()["roles"]) == {"buyer", "broker", "retailer"}


async def test_non_admin_cannot_touch_reminder_settings(fresh_db):
    c = cbq("ar:toggle", 5001)
    await B.admin_reminders(c, FakeState())
    c.answer.assert_awaited_with("فقط برای مدیران.", show_alert=True)


async def test_admin_can_edit_reminder_text(fresh_db):
    admin = settings.admin_chat_ids[0]
    st = FakeState()
    await B.admin_reminders(cbq("ar:text", admin), st)
    await B.save_rem_text(msg("سلام {name}، دوره‌های {periods} مانده است.", admin), st)
    assert "{name}" in rem.get_config()["text"]
    body = rem.render(rem.get_config()["text"], "شرکت الف", "generator", [("هفتگی", "هفته ۱")])
    assert "شرکت الف" in body and "هفته ۱" in body


# ---------- خروجی ادمین ----------

async def test_excel_keeps_deleted_rows_with_change_trail(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await do_trade("offer", 5001, volume=5000, price=80000)
    st = FakeState()
    await B.order_edit_start(cbq("oedit:o:1:volume"), st)
    await B.order_edit_save(msg("7000"), st)
    await B.order_delete(cbq("odelok:o:1"))

    from app.export import build_workbook
    wb = load_workbook(BytesIO(build_workbook()))
    assert set(wb.sheetnames) == {"سازمان‌ها", "عرضه‌ها", "تقاضاها", "تاریخچه تغییرات"}

    ws = wb["عرضه‌ها"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, [c.value for c in ws[2]]))
    assert row["وضعیت"] == "حذف‌شده" and row["تاریخ حذف"]        # رکورد باقی مانده است
    trail = row["تاریخچه تغییرات"]
    assert "۵,۰۰۰ کیلووات" in trail and "۷,۰۰۰ کیلووات" in trail   # مقدار قبلی و جدید هر دو هستند
    assert "حذف" in trail and "ایجاد" in trail and "1405/" in trail
    assert row["نام شخص"] == "علی رضایی" and row["شماره تلفن"] == "09121234567"
    assert row["کف قیمت (ریال)"] == 80000
    assert wb["تاریخچه تغییرات"].max_row >= 4


async def test_report_and_panel_are_admin_only(fresh_db):
    for handler in (B.report, B.sheet, B.stats, B.admin_panel):
        m = msg("/x", 5001)
        await handler(m)
        assert "فقط برای مدیران" in said(m)
    admin = msg("/report", settings.admin_chat_ids[0])
    await B.report(admin)
    admin.answer_document.assert_awaited_once()


async def test_sheet_command_reports_exactly_what_is_missing(fresh_db):
    m = msg("/sheet", settings.admin_chat_ids[0])
    await B.sheet(m)
    out = said(m)
    assert "GOOGLE_SHEET_ID" in out and "Service Account" in out


# ---------- محرمانگی و ثبت پایه ----------

async def test_no_cross_user_data_leak(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await do_trade("offer", 5001, volume=5000, price=80000)
    await register(5002, "buyer", "فولاد نمونه", "3000")
    final, _, _ = await do_trade("request", 5002, volume=4000, price=85000)
    assert "آرمانی" not in said(final) and "تطبیق" not in said(final)

    m = msg(B.BTN_ORDERS, 5002)
    await B.my_orders(m)
    assert "آرمانی" not in said(m) and "فولاد نمونه" in said(m)


async def test_orders_are_numbered_per_profile(fresh_db):
    await register(5001, "generator", "نیروگاه اول")
    await do_trade("offer", 5001)
    await do_trade("offer", 5001, period_index=1)
    await register(5002, "generator", "نیروگاه دوم")
    final, _, _ = await do_trade("offer", 5002)
    conn = connect()
    assert [r["seq"] for r in conn.execute("SELECT seq FROM offers ORDER BY id")] == [1, 2, 1]
    conn.close()
    assert "<b>شماره سفارش:</b> ۱" in said(final)      # شماره‌ی پروفایل، نه شناسه‌ی سراسری ۳


async def test_persian_digits_accepted_for_amounts(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    final, _, _ = await do_trade("offer", 5001, volume="۵۰۰۰", price="۸۰۰۰۰")
    conn = connect()
    row = conn.execute("SELECT * FROM offers ORDER BY id DESC LIMIT 1").fetchone()
    conn.close()
    assert row["volume_kw"] == 5000 and row["min_price"] == 80000


async def test_invalid_amounts_rejected(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    st = FakeState()
    await B.start_trade("offer", msg(B.BTN_OFFER), st)
    await B.choose_period_type(cbq("pt:offer:weekly"), st)
    await B.choose_period(cbq("pd:offer:0"), st)
    for bad_text in ("-3", "abc", "99999999999"):
        bad = msg(bad_text)
        await B.trade_volume(bad, st)
        assert "نامعتبر" in said(bad)


async def test_periods_are_valid_future_jalali(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    code, label, start = cal.available_periods("weekly", calendar_rules())[0]
    _, kb, _ = await do_trade("offer", 5001)
    assert len(kb.inline_keyboard) == calendar_rules().future_periods_visible
    assert kb.inline_keyboard[0][0].text == label and start > jdatetime.date.today()
    conn = connect()
    assert conn.execute("SELECT delivery_period FROM offers").fetchone()[0] == code
    conn.close()


# ---------- راهنما ----------

async def test_help_is_bulleted_and_shown_to_new_users(fresh_db):
    m = msg("/help")
    await B.help_cmd(m)
    out = said(m)
    assert "راهنمای سریع" in out and out.count("▫️") >= 8
    for section in ("ثبت‌نام", "ثبت سفارش", "مدیریت سفارش‌ها", "پروفایل و نقش‌ها"):
        assert f"<b>" in out and section in out
    assert "کیلووات" in out and "ریال" in out and "@Farbod_Papoli" in out

    newcomer = msg("/start", 7777)
    await B.start(newcomer, FakeState())
    assert "راهنمای سریع" in said(newcomer)          # تازه‌واردها خودکار راهنما می‌گیرند


async def test_help_button_is_in_the_user_menu(fresh_db):
    texts = [btn.text for row in B.menu().keyboard for btn in row]
    assert B.BTN_HELP in texts and B.BTN_ABOUT in texts


# ---------- شماره‌ی قابل تماس ----------

@pytest.mark.parametrize("raw,expected", [
    ("09121234567", "+989121234567"), ("9121234567", "+989121234567"),
    ("+989121234567", "+989121234567"),
])
def test_phone_is_rendered_as_callable_international_number(raw, expected):
    assert B.phone_intl(raw) == expected
    assert expected in B.phone_field(raw)


# ---------- فضای اختصاصی ادمین ----------

async def test_admin_panel_opens_its_own_space(fresh_db):
    admin = settings.admin_chat_ids[0]
    m = msg("/admin", admin)
    await B.admin_panel(m)
    kb = m.answer.call_args_list[-1].kwargs["reply_markup"]
    texts = [btn.text for row in kb.keyboard for btn in row]
    assert B.BTN_A_REPORT in texts and B.BTN_A_REMIND in texts and B.BTN_A_EXIT in texts
    assert B.BTN_OFFER not in texts and B.BTN_ORDERS not in texts   # فضای کاربری اینجا نیست
    assert B.in_admin_mode(admin)

    out = msg(B.BTN_A_EXIT, admin)
    await B.admin_exit(out)
    assert not B.in_admin_mode(admin)
    user_kb = out.answer.call_args_list[-1].kwargs["reply_markup"]
    assert B.BTN_OFFER in [b.text for row in user_kb.keyboard for b in row]


async def test_start_leaves_admin_space(fresh_db):
    admin = settings.admin_chat_ids[0]
    await B.admin_panel(msg("/admin", admin))
    await B.start(msg("/start", admin), FakeState())
    assert not B.in_admin_mode(admin)


async def test_non_admin_gets_no_admin_space(fresh_db):
    m = msg("/admin", 5001)
    await B.admin_panel(m)
    assert "فقط برای مدیران" in said(m)
    assert not B.in_admin_mode(5001)


async def test_admin_latest_orders_include_contact_details(fresh_db):
    await register(5001, "generator", "آرمانی انرژی آسیا")
    await do_trade("offer", 5001, volume=5000, price=80000)
    m = msg(B.BTN_A_ORDERS, settings.admin_chat_ids[0])
    await B.latest_orders(m)
    out = said(m)
    assert "آرمانی انرژی آسیا" in out and "علی رضایی" in out
    assert "+989121234567" in out and "<b>کف قیمت:</b>" in out


async def test_admin_handlers_are_registered_on_the_dedicated_dispatcher(fresh_db):
    assert B.admin_dp is not B.dp
    assert len(B.admin_dp.message.handlers) >= 12
    assert B.SEPARATE_ADMIN_BOT is False        # در تست، توکن جدا تنظیم نشده است
    assert B.admin_bot is B.bot


async def test_admin_reminders_panel_lists_current_settings(fresh_db):
    m = msg(B.BTN_A_REMIND, settings.admin_chat_ids[0])
    await B.reminders_panel(m)
    out = said(m)
    assert "یادآوری روزانه" in out and "<b>ساعت‌های ارسال:</b>" in out
    kb = m.answer.call_args_list[-1].kwargs["reply_markup"]
    datas = [x[0].callback_data for x in kb.inline_keyboard]
    assert "ar:toggle" in datas and "ar:times" in datas and "ar:roles" in datas


# ---------- تزریق وابستگی‌های aiogram ----------

async def test_admin_handlers_accept_aiogram_injected_kwargs(fresh_db):
    """aiogram وابستگی‌هایی مثل dispatcher/bot را به هندلر تزریق می‌کند.

    رگرسیون: دکوراتور admin_only نباید این آرگومان‌ها را به تابع اصلی فوروارد کند.
    """
    admin = settings.admin_chat_ids[0]
    injected = dict(dispatcher=object(), bot=B.bot, event_update=object(),
                    event_router=object(), fsm_storage=object(), handler=object())
    for handler in (B.report, B.stats, B.sheet, B.latest_orders, B.reminders_panel):
        m = msg("/x", admin)
        await handler(m, **injected)                     # نباید TypeError بدهد
        assert said(m) or m.answer_document.await_count

    m = msg("/setabout", admin)
    await B.set_about_start(m, FakeState(), **injected)
    assert "درباره ما" in said(m)


def test_admin_only_preserves_the_original_signature():
    """aiogram با inspect.signature تصمیم می‌گیرد چه چیزی تزریق کند."""
    import inspect as _i
    assert list(_i.signature(B.report).parameters) == ["message"]
    assert list(_i.signature(B.set_about_start).parameters) == ["message", "state"]
    for h in (B.report, B.stats, B.sheet, B.latest_orders, B.reminders_panel):
        assert "dispatcher" not in _i.signature(h).parameters
