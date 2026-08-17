"""خروجی داده‌ها برای ادمین — اکسل و Google Sheets از یک منبع مشترک تغذیه می‌شوند.

اصل کار: هیچ رکوردی از گزارش حذف نمی‌شود. ویرایش و حذف فقط وضعیت رکورد را عوض می‌کنند
و ردپای کامل تغییرات در ستون «تاریخچه تغییرات» و شیت جداگانه باقی می‌ماند.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import connect, audit_trail, ROLE_LABELS
from .calendar_utils import period_label

PTYPE = {"weekly": "هفتگی", "monthly": "ماهانه"}
ENTITY = {"organization": "پروفایل", "offer": "عرضه", "request": "تقاضا"}
ACTION = {"create": "ایجاد", "update": "ویرایش", "delete": "حذف"}

ORG_HEADERS = ["شناسه", "نام شرکت", "نقش", "نام شخص", "سمت", "شماره تلفن", "سقف مجاز (kW)",
               "وضعیت", "تاریخ ثبت", "آخرین ویرایش", "تاریخ حذف", "تاریخچه تغییرات",
               "شناسه تلگرام", "یوزرنیم"]
ORDER_HEADERS = ["شناسه سراسری", "شماره سفارش", "شرکت", "نقش", "نام شخص", "شماره تلفن",
                 "کد دوره", "نوع دوره", "دوره‌ی تحویل", "حجم (kW)", "{price}", "وضعیت",
                 "تاریخ ثبت", "آخرین ویرایش", "تاریخ حذف", "تاریخچه تغییرات", "یوزرنیم"]
AUDIT_HEADERS = ["ردیف", "نوع رکورد", "شناسه رکورد", "عملیات", "مورد", "مقدار قبلی",
                 "مقدار جدید", "زمان (شمسی)", "زمان سامانه"]

ORG_Q = ("SELECT o.*, u.telegram_id, u.username FROM organizations o "
         "LEFT JOIN memberships m ON m.org_id=o.id LEFT JOIN users u ON u.id=m.user_id "
         "GROUP BY o.id ORDER BY o.id")
ORDER_Q = ("SELECT t.*, t.{price} AS price, g.name AS org, g.role, g.contact_person, g.phone, "
           "u.username FROM {table} t LEFT JOIN organizations g ON g.id=t.org_id "
           "LEFT JOIN users u ON u.id=t.user_id ORDER BY t.id")


def _status(r):
    return "حذف‌شده" if r["deleted_at"] else "فعال"


def sheet_data():
    """{نام شیت: (سرستون‌ها، سطرها)} — منبع مشترک اکسل و گوگل‌شیت."""
    conn = connect()
    out = {}

    out["سازمان‌ها"] = (ORG_HEADERS, [
        [r["id"], r["name"], ROLE_LABELS.get(r["role"], r["role"]), r["contact_person"],
         r["position"], r["phone"], r["capacity_kw"], _status(r), r["created_at"],
         r["updated_at"], r["deleted_at"], audit_trail(conn, "organization", r["id"]),
         r["telegram_id"], r["username"]]
        for r in conn.execute(ORG_Q).fetchall()])

    for sheet, table, price_col, price_label, entity in (
            ("عرضه‌ها", "offers", "min_price", "کف قیمت (ریال)", "offer"),
            ("تقاضاها", "requests", "max_price", "سقف قیمت (ریال)", "request")):
        rows = conn.execute(ORDER_Q.format(price=price_col, table=table)).fetchall()
        out[sheet] = ([h.replace("{price}", price_label) for h in ORDER_HEADERS], [
            [r["id"], r["seq"], r["org"], ROLE_LABELS.get(r["role"], r["role"]),
             r["contact_person"], r["phone"], r["delivery_period"],
             PTYPE.get(r["period_type"], r["period_type"] or ""),
             period_label(r["delivery_period"], r["period_type"]), r["volume_kw"], r["price"],
             _status(r), r["created_at"], r["updated_at"], r["deleted_at"],
             audit_trail(conn, entity, r["id"]), r["username"]]
            for r in rows])

    out["تاریخچه تغییرات"] = (AUDIT_HEADERS, [
        [r["id"], ENTITY.get(r["entity"], r["entity"]), r["entity_id"],
         ACTION.get(r["action"], r["action"]), r["field"], r["old_value"], r["new_value"],
         r["created_at_jalali"], r["created_at"]]
        for r in conn.execute("SELECT * FROM audit_log ORDER BY id").fetchall()])

    conn.close()
    return out


HEAD_FILL = PatternFill("solid", fgColor="1F4E78")


def build_workbook():
    """خروجی اکسل به‌صورت bytes (برای ارسال در تلگرام و API)."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheet_data().items():
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            width = 45 if "تاریخچه" in h else max(12, len(h) + 6)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
